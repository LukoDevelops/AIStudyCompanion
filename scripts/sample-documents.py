from pathlib import Path
import json
import re
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor

root = Path(__file__).resolve().parents[1] / 'samples' / 'text'
root.mkdir(parents=True, exist_ok=True)
metadata = json.loads(Path(sys.argv[1]).read_text(encoding='utf-8'))
abstract = re.sub(r'(?<=\w)-\s+(?=\w)', '', metadata['abstract'])
abstract = re.sub(r'\s+', ' ', abstract).replace('systems.We', 'systems. We').replace('oper-ations', 'operations').strip()
title = metadata['title']
credit = 'Jeremy D. Frank, NASA Ames Research Center, 2019'
source = 'https://ntrs.nasa.gov/citations/20190032627'
notice = 'Abstract from NASA NTRS. Line-break hyphenation has been repaired. This reading copy is reformatted, not an original publisher Word document.'
(root / 'nasa-ai-abstract.md').write_text(f'# {title}\n\n{credit}\n\n## Abstract\n\n{abstract}\n\n## Source\n\n{source}\n\n{notice}\n', encoding='utf-8')
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.8)
section.left_margin = section.right_margin = Inches(.8)
for name in ['Normal', 'Title', 'Heading 1']:
    doc.styles[name].font.name = 'Calibri'
    doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(8)
doc.styles['Title'].font.size = Pt(20)
doc.add_paragraph(title, 'Title')
doc.add_paragraph(credit)
doc.add_heading('Abstract', 1)
doc.add_paragraph(abstract)
doc.add_heading('Source', 1)
doc.add_paragraph(source)
doc.add_paragraph(notice)
doc.core_properties.title = title
doc.core_properties.author = 'NASA Ames Research Center'
doc.core_properties.comments = 'Reformatted abstract; source and conversion notice are included in the document.'
doc.save(root / 'nasa-ai-abstract.docx')
print('Prepared NASA abstract in Markdown and Word.')
